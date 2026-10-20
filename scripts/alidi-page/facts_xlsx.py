#!/usr/bin/env python3
"""Таблица фактов и цифр для АЛИДИ: что знаем из ТЗ, alidi.ru и встречи 7 октября, клиент дополняет.

    python3 scripts/alidi-page/facts_xlsx.py   # ~/Documents/Тендеры/АЛИДИ/Hand_Marketing_АЛИДИ_факты_2026-10-20.xlsx

Обещана в протоколе на странице /for/alidi/ («Таблица фактов и цифр для заполнения»).
"""
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

OUT = Path.home() / 'Documents' / 'Тендеры' / 'АЛИДИ' / 'Hand_Marketing_АЛИДИ_факты_2026-10-20.xlsx'

# (раздел, [(показатель, что знаем сейчас, откуда)])
DATA = [
    ('Компания', [
        ('Год и город основания', '1992, Нижний Новгород', 'alidi.ru, встреча 7.10'),
        ('Сотрудников всего', '9 700 в ТЗ, более 10 000 на сайте: нужно одно число', 'ТЗ, alidi.ru'),
        ('Сотрудников по странам: Россия, Беларусь, Казахстан', '', ''),
        ('Оборот', 'около 100 млрд ₽ на сайте: за какой период?', 'alidi.ru'),
        ('Брендов в портфеле', 'более 100', 'alidi.ru'),
        ('Торговых точек, с которыми работаете', 'более 100 тыс.', 'alidi.ru'),
        ('Регионов присутствия', 'Россия 62, Беларусь 7, Казахстан 9', 'alidi.ru/logistic'),
        ('Каналы и направления', 'розница, HoReCa, АЗС, e-com, «На_полке»', 'история на alidi.ru'),
    ]),
    ('Склады', [
        ('Россия: складов и площадь', '15 складов, 143 000 м²', 'alidi.ru/logistic'),
        ('Беларусь: складов и площадь', '1 склад, 19 000 м²', 'alidi.ru/logistic'),
        ('Казахстан: складов и площадь', '12 складов, 31 000 м²', 'alidi.ru/logistic'),
        ('Всего складов и площадь', '28 складов, 193 000 м²: подтвердите сумму', 'сумма по сайту'),
        ('Валищево: площадь, температурные зоны, высота хранения', '', ''),
        ('Нижний Новгород: площадь склада', '', ''),
        ('Санкт-Петербург: площадь склада', '', ''),
        ('Минск: площадь склада', '', ''),
        ('Алматы: площадь склада', '', ''),
        ('Казань (опция): площадь склада', '', ''),
        ('Отгрузок, заказов или паллет в сутки', '', ''),
    ]),
    ('Площадки съёмки', [
        ('Москва, офис: адрес, контакт на площадке', '', ''),
        ('Валищево: начало первой и конец последней смены', 'ночные смены есть, снимать можно по согласованию', 'встреча 7.10'),
        ('Нижний Новгород: время смен, контакт', 'склад и офис в одном месте', 'встреча 7.10'),
        ('Санкт-Петербург: время смен, контакт', '', ''),
        ('Минск: время смен, контакт', '', ''),
        ('Алматы: время первой смены по местному времени, контакт', '', ''),
        ('Дни с наименьшей нагрузкой на складах', 'зону склада можно остановить на 1–3 часа', 'встреча 7.10'),
    ]),
    ('Люди', [
        ('Генеральный директор', 'Иван Сычёв, 23 года в компании, начинал в Нижнем Новгороде', 'встреча 7.10'),
        ('Руководители для интервью: ФИО и должность', '', ''),
        ('Сотрудники с самым большим стажем: ФИО, год прихода, площадка', '', ''),
        ('Сотрудники, пришедшие недавно (2026–2027)', '', ''),
    ]),
    ('История и архив', [
        ('Вехи для фильма', '1993 P&G, 2001 Nestlé, 2005 логистический центр в Нижнем, 2007 Wrigley и Purina, 2010 Mars, '
         '2011 Беларусь, 2016 Казахстан и HoReCa, 2017 Nestlé Professional, 2022 АЗС, 2023 e-com и «На_полке»', 'история на alidi.ru'),
        ('Центральный офис по годам', 'Нижний Новгород, затем Санкт-Петербург, около 2019–2021 Москва', 'встреча 7.10'),
        ('Архивные фото и видео: что есть, с какого года', 'ролик «Алиди 25 лет» на YouTube', 'встреча 7.10'),
    ]),
    ('Бренды и награды', [
        ('Партнёры, которых можно показывать и называть', 'список пришлёте вы', 'встреча 7.10'),
        ('Логотипы, которые в кадре нельзя', '', ''),
        ('Награды и рейтинги: название и год (например, ТОП-100)', '', ''),
    ]),
    ('Показ фильма', [
        ('Дата и площадка мероприятия к 35-летию', '', ''),
        ('Экран в зале: размер, формат, звук', '', ''),
        ('Где ещё покажете: сайт, HH.ru, тендеры, соцсети', '', ''),
    ]),
]


def build(path=OUT):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment
    from openpyxl.worksheet.datavalidation import DataValidation
    import xlsx_brand as X

    S = X.styles()
    font, fill = S['font'], S['fill']
    wb = Workbook()
    ws = wb.active
    ws.title = 'Факты'
    colw = (6, 44, 50, 20, 34, 14, 14, 3)
    for i, w in enumerate(colw):
        ws.column_dimensions[chr(65 + i)].width = w
    W = sum(X.col_px(w) for w in colw)
    n_rows = 8 + sum(len(r) + 2 for _, r in DATA) + 6
    X.paper(ws, n_rows, len(colw))
    fams = [X.FAMILY[k] for k in X.ORDER]
    X.render_figs(sorted({f[3] for f in fams} | {'fig_06', 'fig_01'}), 40)
    out = X.CACHE / 'banner-facts.png'
    X.render([{'out': str(out), 'w': W, 'h': 240, 'html': X.banner_html(
        W, 240, 'ГК АЛИДИ · имиджевый фильм к 35-летию', 'Факты и цифры для фильма',
        'Hand Marketing · 20 октября 2026 · в кадр пойдут только цифры, которые вы подтвердите', ['fig_01', 'fig_06', 'fig_15'])}])
    ws.row_dimensions[1].height = 240 * 0.75
    X.put_image(ws, out, 'A1', W, 240)

    ws['B3'] = ('Во второй колонке то, что мы нашли в ТЗ, на alidi.ru и услышали на встрече 7 октября. Жёлтые колонки ваши: '
                'поправьте или впишите значение, дату, на которую оно верно, и можно ли показывать его в фильме. '
                'Заполняйте, что удобно, остальное уточним по ходу.')
    ws.merge_cells('B3:G3'); ws['B3'].font = font(10, color=X.STONE[600]); ws['B3'].alignment = S['wrap']
    ws.row_dimensions[3].height = 42
    ws.merge_cells('B5:D5'); ws['B5'] = 'ЧТО ЗНАЕМ МЫ'
    ws['B5'].font = font(8, True, X.STONE[600]); ws['B5'].fill = fill(X.STONE[100]); ws['B5'].alignment = S['mid_c']
    ws.merge_cells('E5:G5'); ws['E5'] = 'ЗАПОЛНЯЕТЕ ВЫ'
    ws['E5'].font = font(8, True, X.INK); ws['E5'].fill = fill(X.INPUT_FILL); ws['E5'].alignment = S['mid_c']
    ws.row_dimensions[5].height = 18
    for i, t in enumerate(['Показатель', 'Что знаем сейчас', 'Откуда', 'Ваше значение', 'На какую дату', 'В фильм можно?']):
        c = ws.cell(6, 2 + i, t); c.font = font(8, True, X.PAPER); c.fill = fill(X.INK); c.alignment = S['mid']
    ws.row_dimensions[6].height = 24
    dv = DataValidation(type='list', formula1='"да,нет,уточним"', allow_blank=True)
    ws.add_data_validation(dv)
    r = 7
    for si, (sec, rows) in enumerate(DATA):
        main, light, tfg, fig = fams[si % len(fams)]
        r += 1
        X.put_image(ws, X.fig_png(fig), f'A{r}', 34, 34)
        for k in range(2, 8):
            ws.cell(r, k).fill = fill(main)
        ws.cell(r, 2, f'{si + 1:02d}   {sec}').font = font(12, True, tfg)
        ws.cell(r, 2).alignment = S['mid']
        ws.row_dimensions[r].height = 28
        r += 1
        for name, known, src in rows:
            ws.cell(r, 2, name).font = font(10, True)
            ws.cell(r, 3, known or None).font = font(9, color=X.STONE[600])
            ws.cell(r, 4, src or None).font = font(8, color=X.STONE[400])
            for k in (5, 6, 7):
                ws.cell(r, k).fill = fill(X.INPUT_FILL)
                ws.cell(r, k).font = font(10)
            dv.add(f'G{r}')
            for k in range(2, 8):
                ws.cell(r, k).border = S['line']
                ws.cell(r, k).alignment = Alignment(vertical='center', wrap_text=True)
            ln = max(len(name) / 50, len(known or '') / 55)
            ws.row_dimensions[r].height = max(24, 13 * (int(ln) + 1) + 8)
            r += 1
    X.footer(ws, r + 1, 7, S)
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = X.FAMILY['cyan'][0]
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    wb.properties.creator = wb.properties.lastModifiedBy = 'Hand Marketing'
    wb.properties.title = 'ГК АЛИДИ · имиджевый фильм к 35-летию · факты и цифры'
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


if __name__ == '__main__':
    print(build(Path(sys.argv[1]) if len(sys.argv) > 1 else OUT))
